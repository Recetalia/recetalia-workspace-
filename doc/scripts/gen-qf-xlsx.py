"""Genera el Excel de QF por farmacia (datos de PROD) para el relevamiento de emails."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FARMACIAS = json.load(open("/tmp/qf_farmacias.json"))
QFS = json.load(open("/tmp/qf_consolidado.json"))
SALIDA = "/Users/pablo/Downloads/QF-por-farmacia-PROD-2026-08-13.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
COMPLETAR_FILL = PatternFill("solid", fgColor="FFF2CC")  # amarillo: a completar
REVISAR_FILL = PatternFill("solid", fgColor="FCE4E4")    # rojo suave: CJP a revisar


def limpio(v):
    """MySQL devuelve el string 'null' cuando el JSON no traía el campo."""
    if v is None or str(v).strip().lower() in ("null", "none"):
        return ""
    return str(v).strip()


def telefono(raw):
    """pharmacy.phone es un JSON: {countryCode, national, international, ...}."""
    raw = limpio(raw)
    if not raw.startswith("{"):
        return raw
    try:
        d = json.loads(raw)
    except json.JSONDecodeError:
        return raw
    return limpio(d.get("international") or d.get("national") or "")


def escribir_encabezado(ws, columnas):
    ws.append([c[0] for c in columnas])
    for i, (_, ancho) in enumerate(columnas, start=1):
        celda = ws.cell(row=1, column=i)
        celda.fill, celda.font = HEADER_FILL, HEADER_FONT
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 30


# ---------------------------------------------------------------- Hoja 1
wb = Workbook()
ws = wb.active
ws.title = "Farmacias"
columnas = [
    ("Farmacia", 30), ("Cadena", 14), ("Localidad", 18),
    ("Email de la farmacia", 32), ("Teléfono farmacia", 18),
    ("QF asignado", 28), ("CJP", 10), ("Documento QF", 14), ("Estado QF", 14),
    ("D.T. declarado por la farmacia", 28), ("CJP declarado", 12),
    ("➜ EMAIL DEL QF (completar)", 32), ("➜ Observaciones", 26),
]
escribir_encabezado(ws, columnas)

# QF de prueba que quedaron en la base de PROD (medido 2026-08-13): no hay que contactarlos.
CJP_PRUEBA = {"446788", "453652", "9876543"}


def es_prueba(cjp):
    return limpio(cjp) in CJP_PRUEBA


# Las de prueba al final, para que no estorben el trabajo de contacto.
filas = sorted(FARMACIAS, key=lambda f: (es_prueba(f["qf_cjp"]), limpio(f["cadena"]), limpio(f["farmacia"])))
for f in filas:
    revisar = limpio(f["qf_estado"]) == "NEEDS_REVIEW"
    prueba = es_prueba(f["qf_cjp"])
    ws.append([
        limpio(f["farmacia"]), limpio(f["cadena"]) or "(independiente)", limpio(f["localidad"]),
        limpio(f["email_farmacia"]), telefono(f["tel_farmacia"]),
        limpio(f["qf_nombre"]), limpio(f["qf_cjp"]), limpio(f["qf_doc"]),
        "CJP A REVISAR" if revisar else "OK",
        limpio(f["dt_declarado"]), limpio(f["cjp_declarado"]),
        "", "REGISTRO DE PRUEBA — no contactar" if prueba else "",
    ])
    fila = ws.max_row
    ws.cell(row=fila, column=12).fill = COMPLETAR_FILL
    ws.cell(row=fila, column=13).fill = COMPLETAR_FILL
    if revisar:
        ws.cell(row=fila, column=9).fill = REVISAR_FILL
    if prueba:
        ws.cell(row=fila, column=13).fill = REVISAR_FILL
        ws.cell(row=fila, column=13).font = Font(bold=True)
ws.auto_filter.ref = f"A1:M{ws.max_row}"

# ---------------------------------------------------------------- Hoja 2
ws2 = wb.create_sheet("QF consolidado")
columnas2 = [
    ("CJP", 10), ("Químico Farmacéutico", 30), ("Documento", 14), ("Estado", 14),
    ("Cant. farmacias", 14), ("Farmacias a su cargo", 70), ("➜ EMAIL DEL QF (completar)", 32),
]
escribir_encabezado(ws2, columnas2)
for d in sorted(QFS, key=lambda x: (es_prueba(x["cjp"]), -(x["cant_farmacias"] or 0))):
    revisar = limpio(d["estado"]) == "NEEDS_REVIEW"
    ws2.append([
        limpio(d["cjp"]), limpio(d["nombre"]), limpio(d["doc"]),
        "PRUEBA — no contactar" if es_prueba(d["cjp"]) else ("CJP A REVISAR" if revisar else "OK"),
        d["cant_farmacias"] or 0, limpio(d["farmacias"]), "",
    ])
    fila = ws2.max_row
    ws2.cell(row=fila, column=7).fill = COMPLETAR_FILL
    if revisar:
        ws2.cell(row=fila, column=4).fill = REVISAR_FILL
ws2.auto_filter.ref = f"A1:G{ws2.max_row}"

# ---------------------------------------------------------------- Hoja 3
ws3 = wb.create_sheet("Resumen")
ws3.column_dimensions["A"].width = 52
ws3.column_dimensions["B"].width = 14
sin_email = sum(1 for d in QFS if not limpio(d.get("email", "")))
con_doc = sum(1 for d in QFS if limpio(d["doc"]))
revisar_qf = sum(1 for d in QFS if limpio(d["estado"]) == "NEEDS_REVIEW")
revisar_farm = sum(1 for f in FARMACIAS if limpio(f["qf_estado"]) == "NEEDS_REVIEW")
prueba_farm = sum(1 for f in FARMACIAS if es_prueba(f["qf_cjp"]))
sin_email_farm = sum(1 for f in FARMACIAS if not limpio(f["email_farmacia"]))
for texto, valor in [
    ("Fuente: DB de PRODUCCIÓN (recetali_receta), 2026-08-13", ""),
    ("", ""),
    ("Farmacias activas", len(FARMACIAS)),
    ("  de las cuales son registros de prueba", prueba_farm),
    ("  FARMACIAS REALES A CONTACTAR", len(FARMACIAS) - prueba_farm),
    ("Farmacias sin QF asignado", 0),
    ("Farmacias sin email propio (no se las puede contactar)", sin_email_farm),
    ("", ""),
    ("Químicos Farmacéuticos", len(QFS)),
    ("QF SIN email cargado", sin_email),
    ("QF con email cargado", len(QFS) - sin_email),
    ("QF con teléfono cargado", 0),
    ("QF con cédula cargada (el resto viene {number:null})", con_doc),
    ("", ""),
    ("QF con CJP a revisar (NEEDS_REVIEW)", revisar_qf),
    ("Farmacias afectadas por esos CJP", revisar_farm),
]:
    ws3.append([texto, valor])
ws3["A1"].font = Font(bold=True)
for fila in (3, 7, 11):
    ws3.cell(row=fila, column=1).font = Font(bold=True)

wb.save(SALIDA)
print("OK:", SALIDA)
print(f"Hoja Farmacias: {ws.max_row - 1} filas | QF consolidado: {ws2.max_row - 1} filas")
print(f"QF sin email: {sin_email}/{len(QFS)} | NEEDS_REVIEW: {revisar_qf} QF / {revisar_farm} farmacias")
